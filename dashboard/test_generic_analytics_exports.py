from copy import deepcopy
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from dashboard.selectors.analytics import (
    get_order_metrics,
    get_product_metrics,
    resolve_analytics_period,
)
from dashboard.selectors.analytics_exports import build_xlsx
from dashboard.selectors.catalog_reporting import (
    LEGACY_UNKNOWN_UNIT,
    project_current_inventory,
    project_purchased_lines,
)
from orders.models import Order, OrderItem
from products.catalog_domain import serialize_combination_signature
from products.models import (
    AttributeDefinition,
    AttributeOption,
    Product,
    ProductAttributeValue,
    ProductType,
    ProductTypeAttribute,
    ProductVariant,
    VariantAttributeValue,
)


def _attribute(code, label, scope, value_code, display, order, *, dimension=False):
    return {
        "definition_code": code,
        "label_default": label,
        "label_fa": label,
        "scope": scope,
        "data_type": "option",
        "unit": None,
        "is_dimension": dimension,
        "sort_order": order,
        "values": [{
            "option_code": value_code,
            "raw_value": value_code,
            "display_default": display,
            "display_fa": display,
        }],
    }


def _v2_snapshot(*, product_id=101, variant_id=201, unit="unit", product_name="iPhone 17 Pro"):
    dimensions = [
        _attribute("color", "Color", "variant", "black", "Black", 10, dimension=True),
        _attribute("ram", "RAM", "variant", "12gb", "12 GB", 20, dimension=True),
        _attribute("storage", "Storage", "variant", "256gb", "256 GB", 30, dimension=True),
    ]
    return {
        "version": 2,
        "product": {
            "id": product_id,
            "product_type": {"code": "mobile", "label_default": "Mobile", "label_fa": "Mobile"},
            "sku": "IPHONE-17", "name_default": product_name, "name_fa": product_name,
        },
        "variant": {
            "id": variant_id, "sku": "IP17-BLK-12-256",
            "combination_signature": "color=black|ram=12gb|storage=256gb",
            "selling_unit": {
                "code": unit, "price_label_default": "Per Unit", "price_label_fa": "Per Unit",
                "singular_noun_default": "item", "singular_noun_fa": "item",
                "plural_noun_default": "items", "plural_noun_fa": "items",
            },
        },
        "attributes": {
            "product": [_attribute("brand", "Brand", "product", "apple", "Apple", 10)],
            "variant": dimensions,
        },
        "operational": {},
    }


class GenericHistoricalReportingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user("report-buyer")

    def _order(self, currency="USD"):
        return Order.objects.create(
            user=self.user, full_name="Buyer", address="Address", city="City",
            phone="123", currency=currency, total_amount=Decimal("100.00"),
            is_paid=True, paid_at=timezone.now(),
        )

    def _item(self, snapshot, *, quantity=2, price="10.00", product_name="Frozen name"):
        return OrderItem.objects.create(
            order=self._order(), product_name=product_name,
            quantity=quantity, unit_price=Decimal(price), variant_snapshot=snapshot,
        )

    def test_v2_projection_uses_frozen_generic_identity_and_attributes(self):
        item = self._item(_v2_snapshot())
        line = project_purchased_lines([item])[0]
        self.assertEqual(line.product_name, "iPhone 17 Pro")
        self.assertEqual(line.selling_unit_code, "unit")
        self.assertEqual([row.definition_code for row in line.variant_attributes], ["color", "ram", "storage"])
        self.assertEqual(line.variant_attributes[1].stable_values, ("12gb",))
        self.assertEqual(line.revenue, Decimal("20.00"))

    def test_legacy_snapshot_unit_is_explicitly_unknown_even_with_live_variant(self):
        product = Product.objects.create(name="Current", slug="current", sku="CURRENT")
        variant = ProductVariant.objects.create(product=product, sku="CURRENT-V", price_unit="box")
        item = OrderItem.objects.create(
            order=self._order(), product=product, variant=variant,
            product_name="Frozen Tile", quantity=3, unit_price=Decimal("4.00"),
            variant_snapshot={"sku": "OLD-V", "size": "60x60", "color": "Beige"},
        )
        line = project_purchased_lines([item])[0]
        self.assertEqual(line.product_name, "Frozen Tile")
        self.assertEqual(line.variant_sku, "OLD-V")
        self.assertEqual(line.selling_unit_code, LEGACY_UNKNOWN_UNIT)

    def test_projection_is_immune_to_live_catalog_mutation_and_deletion(self):
        frozen = _v2_snapshot()
        item = self._item(deepcopy(frozen))
        first = project_purchased_lines([item])[0]
        item.product_name = "Mutated OrderItem fallback"
        item.save(update_fields=["product_name"])
        second = project_purchased_lines([item])[0]
        self.assertEqual(first, second)
        self.assertEqual(second.product_name, "iPhone 17 Pro")

    def test_malformed_snapshot_is_diagnostic_and_never_healed(self):
        item = self._item({"version": 2, "product": {}})
        line = project_purchased_lines([item])[0]
        self.assertEqual(line.snapshot_source, "invalid")
        self.assertTrue(line.snapshot_issue)
        self.assertEqual(line.selling_unit_code, LEGACY_UNKNOWN_UNIT)

    def test_mixed_selling_units_are_grouped_not_summed(self):
        self._item(_v2_snapshot(variant_id=1, unit="unit"), quantity=2)
        self._item(_v2_snapshot(variant_id=2, unit="box"), quantity=3)
        self._item({"sku": "LEGACY", "size": "60x60"}, quantity=4)
        metrics = get_order_metrics(resolve_analytics_period({"range": "today"}))
        self.assertEqual(metrics["quantity_sold_by_unit"], {
            "box": 3, LEGACY_UNKNOWN_UNIT: 4, "unit": 2,
        })
        self.assertNotIn("units_sold", metrics)

    def test_top_products_are_snapshot_authoritative_and_currency_bound(self):
        self._item(_v2_snapshot(product_name="Frozen Mobile"), quantity=2, price="12.50")
        row = get_product_metrics(resolve_analytics_period({"range": "today"}))["top_products"][0]
        self.assertEqual(row["product_name"], "Frozen Mobile")
        self.assertEqual(row["currency"], "USD")
        self.assertEqual(row["quantity_by_unit"], {"unit": 2})
        self.assertEqual(row["net_revenue"], Decimal("25.00"))

    def test_xlsx_uses_dynamic_scoped_headers_and_no_fixed_tile_columns(self):
        self._item(_v2_snapshot(), quantity=2)
        workbook = load_workbook(BytesIO(build_xlsx(resolve_analytics_period({"range": "today"}))))
        headers = [cell.value for cell in workbook["Variants"][1]]
        self.assertIn("Product • Brand [brand]", headers)
        self.assertIn("Variant • RAM [ram]", headers)
        self.assertIn("Variant • Storage [storage]", headers)
        self.assertNotIn("Size", headers)
        self.assertNotIn("Color", headers)
        self.assertNotIn("Finish", headers)
        self.assertIn("Selling Unit", headers)

    def test_multiple_and_scalar_snapshot_values_export_without_loss(self):
        snapshot = _v2_snapshot()
        snapshot["attributes"]["product"].append({
            "definition_code": "features", "label_default": "Features",
            "label_fa": "Features", "scope": "product", "data_type": "option",
            "unit": None, "is_dimension": False, "sort_order": 20,
            "values": [
                {"option_code": "wifi", "raw_value": "wifi", "display_default": "Wi-Fi", "display_fa": "Wi-Fi"},
                {"option_code": "nfc", "raw_value": "nfc", "display_default": "NFC", "display_fa": "NFC"},
            ],
        })
        snapshot["attributes"]["variant"].append({
            "definition_code": "refurbished", "label_default": "Refurbished",
            "label_fa": "Refurbished", "scope": "variant", "data_type": "boolean",
            "unit": None, "is_dimension": False, "sort_order": 40,
            "values": [{"option_code": None, "raw_value": False, "display_default": "No", "display_fa": "No"}],
        })
        self._item(snapshot)
        workbook = load_workbook(BytesIO(build_xlsx(resolve_analytics_period({"range": "today"}))))
        sheet = workbook["Variants"]
        values = dict(zip((cell.value for cell in sheet[1]), (cell.value for cell in sheet[2])))
        self.assertEqual(values["Product • Features [features]"], "Wi-Fi; NFC")
        self.assertEqual(values["Variant • Refurbished [refurbished]"], "No")

    def test_laptop_dimensions_use_the_same_reporting_path(self):
        snapshot = _v2_snapshot(product_name="Framework Laptop")
        snapshot["product"]["product_type"]["code"] = "laptop"
        snapshot["variant"]["sku"] = "LAPTOP-32-1TB"
        snapshot["attributes"]["variant"] = [
            _attribute("color", "Color", "variant", "gray", "Gray", 10, dimension=True),
            _attribute("ram", "RAM", "variant", "32gb", "32 GB", 20, dimension=True),
            _attribute("ssd", "SSD", "variant", "1tb", "1 TB", 30, dimension=True),
        ]
        item = self._item(snapshot)
        line = project_purchased_lines([item])[0]
        self.assertEqual(line.product_type_code, "laptop")
        self.assertEqual([row.definition_code for row in line.variant_attributes], ["color", "ram", "ssd"])

    def test_snapshot_projection_has_zero_queries_after_loading(self):
        item = self._item(_v2_snapshot())
        item = OrderItem.objects.select_related("order").get(pk=item.pk)
        with self.assertNumQueries(0):
            project_purchased_lines([item])


class CurrentGenericInventoryReportingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.mobile = ProductType.objects.create(code="mobile", name="Mobile")
        brand = AttributeDefinition.objects.create(code="brand", name="Brand", data_type="option")
        ram = AttributeDefinition.objects.create(code="ram", name="RAM", data_type="option")
        brand_binding = ProductTypeAttribute.objects.create(
            product_type=cls.mobile, definition=brand, scope="product", sort_order=10,
        )
        cls.ram_binding = ProductTypeAttribute.objects.create(
            product_type=cls.mobile, definition=ram, scope="variant",
            is_variant_dimension=True, is_required=True,
            include_in_purchase_snapshot=True, sort_order=10,
        )
        apple = AttributeOption.objects.create(definition=brand, code="apple", name="Apple")
        cls.ram12 = AttributeOption.objects.create(definition=ram, code="12gb", name="12 GB")
        cls.product = Product.objects.create(
            product_type=cls.mobile, name="Current Mobile", slug="current-mobile",
            sku="MOBILE", is_active=True,
        )
        ProductAttributeValue.objects.create(product=cls.product, binding=brand_binding, option_value=apple)
        cls.variant = ProductVariant.objects.create(
            product=cls.product, sku="MOBILE-12", price_unit="unit", stock=10,
            low_stock_threshold=2, is_active=True,
            combination_signature=serialize_combination_signature((("ram", "12gb"),)),
        )
        VariantAttributeValue.objects.create(variant=cls.variant, binding=cls.ram_binding, option_value=cls.ram12)

    def test_inventory_projection_uses_current_generic_attributes_and_unit(self):
        row = project_current_inventory()[0]
        self.assertEqual(row.product_type_code, "mobile")
        self.assertEqual(row.selling_unit_code, "unit")
        self.assertEqual(row.stock, 10)
        self.assertEqual(row.product_attributes[0].definition_code, "brand")
        self.assertEqual(row.variant_attributes[0].stable_values, ("12gb",))

    def test_inventory_projection_is_batched_for_one_or_many_variants(self):
        with self.assertNumQueries(8):
            project_current_inventory()
        for index in range(3):
            product = Product.objects.create(
                product_type=self.mobile, name=f"Mobile {index}",
                slug=f"mobile-{index}", sku=f"MOBILE-{index}", is_active=True,
            )
            variant = ProductVariant.objects.create(
                product=product, sku=f"COMPLETE-{index}", price_unit="unit",
                stock=index, is_active=True,
                combination_signature=serialize_combination_signature((("ram", "12gb"),)),
            )
            VariantAttributeValue.objects.create(
                variant=variant, binding=self.ram_binding, option_value=self.ram12,
            )
        with self.assertNumQueries(8):
            rows = project_current_inventory()
        self.assertEqual(len(rows), 4)
