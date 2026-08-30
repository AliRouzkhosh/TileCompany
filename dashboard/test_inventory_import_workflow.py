from decimal import Decimal
from io import BytesIO
from tempfile import TemporaryDirectory

from django.conf import settings

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from products.currency import Currency
from products.models import InventoryPriceImportBatch, Product, ProductVariant, VariantPrice
from products.storage import private_inventory_import_storage


class DashboardInventoryImportWorkflowTests(TestCase):
    def setUp(self):
        self.private_storage_directory = TemporaryDirectory()
        self.previous_private_storage_location = private_inventory_import_storage.location
        private_inventory_import_storage.location = self.private_storage_directory.name
        self.addCleanup(self._restore_private_storage)
        call_command("seed_dashboard_roles", verbosity=0)
        User = get_user_model()
        self.manager = User.objects.create_user("inventory-manager", password="secret")
        self.manager.groups.add(Group.objects.get(name="Product Manager"))
        self.run_only = User.objects.create_user("run-only", password="secret")
        self.other = User.objects.create_user("other", password="secret")
        self.run_only.user_permissions.add(Permission.objects.get(content_type__app_label="products", codename="run_inventory_price_import"))
        product = Product.objects.create(name="Import product", slug="import-product", sku="IMPORT")
        self.variant = ProductVariant.objects.create(product=product, sku="IMPORT-1", stock=5, low_stock_threshold=2)
        VariantPrice.objects.create(variant=self.variant, currency=Currency.TOMAN, amount=Decimal("100"))

    def _restore_private_storage(self):
        private_inventory_import_storage.location = self.previous_private_storage_location
        self.private_storage_directory.cleanup()

    def test_permissions_upload_preview_apply_and_private_source_download(self):
        self.client.force_login(self.manager)
        landing = reverse("dashboard:products:inventory_import")
        self.assertContains(self.client.get(landing), "Bulk inventory")
        upload = self.client.post(reverse("dashboard:products:inventory_import_upload"), {
            "source": SimpleUploadedFile("stock.csv", b"sku,stock,low_stock_threshold,irt_price,irt_compare_at,usd_price,usd_compare_at\nIMPORT-1,7,,,,,\n")
        })
        self.assertEqual(upload.status_code, 302)
        batch = InventoryPriceImportBatch.objects.get()
        detail = self.client.get(reverse("dashboard:products:inventory_import_detail", args=[batch.pk]))
        self.assertContains(detail, "IMPORT-1")
        self.assertContains(detail, "Apply 1 change")
        self.assertEqual(self.client.get(reverse("dashboard:products:inventory_import_source_download", args=[batch.pk])).status_code, 200)

        self.client.force_login(self.run_only)
        self.assertEqual(self.client.get(landing).status_code, 200)
        self.assertNotContains(self.client.get(reverse("dashboard:products:inventory_import_detail", args=[batch.pk])), "Apply 1 change")
        self.assertEqual(self.client.post(reverse("dashboard:products:inventory_import_apply", args=[batch.pk])).status_code, 403)

    def test_current_workbook_roundtrips_unchanged_and_is_literal_safe(self):
        self.variant.sku = "=SAFE"; self.variant.save(update_fields=["sku"])
        apostrophe = ProductVariant.objects.create(product=self.variant.product, sku="'SAFE", stock=3, low_stock_threshold=1)
        self.client.force_login(self.manager)
        response = self.client.get(reverse("dashboard:products:inventory_import_current_xlsx"))
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        data = workbook["Data"]
        self.assertEqual(data.cell(2, 1).data_type, "s")
        self.assertEqual(data.cell(2, 1).value, "'=SAFE")
        self.assertIn("'SAFE", [data.cell(row, 1).value for row in range(2, data.max_row + 1)])
        upload = self.client.post(reverse("dashboard:products:inventory_import_upload"), {"source": SimpleUploadedFile("current.xlsx", response.content)})
        batch = InventoryPriceImportBatch.objects.latest("pk")
        self.assertEqual(upload.status_code, 302)
        self.assertEqual(batch.changed_rows, 0)

    def test_templates_and_csv_are_private_normal_navigation_downloads(self):
        self.client.force_login(self.manager)
        self.assertTrue(self.client.get(reverse("dashboard:products:inventory_import_template_csv")).content.startswith(b"\xef\xbb\xbfsku,"))
        self.assertEqual(self.client.get(reverse("dashboard:products:inventory_import_template_xlsx"))["Cache-Control"], "private, no-store")

    def test_source_storage_and_importer_links_are_private_and_full_document(self):
        self.client.force_login(self.manager)
        self.client.post(reverse("dashboard:products:inventory_import_upload"), {
            "source": SimpleUploadedFile("source.csv", b"sku,stock,low_stock_threshold,irt_price,irt_compare_at,usd_price,usd_compare_at\nIMPORT-1,6,,,,,\n")
        })
        batch = InventoryPriceImportBatch.objects.latest("pk")
        self.assertNotIn(str(settings.MEDIA_ROOT), str(batch.source_file.storage.location))
        detail = self.client.get(reverse("dashboard:products:inventory_import_detail", args=[batch.pk]))
        self.assertNotContains(detail, "hx-boost=")
        self.assertNotContains(detail, batch.source_file.name)
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(reverse("dashboard:products:inventory_import_source_download", args=[batch.pk])).status_code, 403)
